"""Tests for ingest readers: dispatch, formats, zip/directory expansion."""

from pathlib import Path
import zipfile

import pandas as pd
import pytest

from knowledge_graph_foundry.ingest import (
    STRUCTURED_EXTENSIONS,
    UNSTRUCTURED_EXTENSIONS,
    UnsupportedFormatError,
    is_structured,
    is_supported,
    iter_source_files,
    read_document,
    read_structured,
)

# -- unstructured -----------------------------------------------------------


def test_read_md(tmp_path):
    path = tmp_path / "note.md"
    path.write_text("# Title\n\nSome content.")
    doc = read_document(path)
    assert doc.text == "# Title\n\nSome content."
    assert doc.format == "md"
    assert doc.id.startswith("d_")
    assert doc.metadata["file_name"] == "note.md"
    assert doc.metadata["file_size"] == path.stat().st_size


def test_read_txt(tmp_path):
    path = tmp_path / "plain.txt"
    path.write_text("plain text body")
    doc = read_document(path)
    assert doc.text == "plain text body"
    assert doc.format == "txt"


def test_read_html(tmp_path):
    path = tmp_path / "page.html"
    path.write_text("<html><body><h1>Head</h1><p>Body text</p></body></html>")
    doc = read_document(path)
    assert "Head" in doc.text
    assert "Body text" in doc.text
    assert "<p>" not in doc.text
    assert doc.format == "html"


def test_read_docx(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("First paragraph")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "cell a"
    table.rows[0].cells[1].text = "cell b"
    path = tmp_path / "doc.docx"
    d.save(str(path))

    doc = read_document(path)
    assert "First paragraph" in doc.text
    assert "cell a | cell b" in doc.text
    assert doc.format == "docx"


def test_read_pdf_mocked(tmp_path, monkeypatch):
    path = tmp_path / "manual.pdf"
    path.write_bytes(b"%PDF-fake")
    monkeypatch.setattr("pymupdf4llm.to_markdown", lambda p: "# PDF markdown")
    # union partner (pypdf) fails on the fake bytes and is tolerated
    doc = read_document(path)
    assert doc.text == "# PDF markdown"
    assert doc.format == "pdf"


def test_pdf_parser_union_appends_pypdf_layer(tmp_path, monkeypatch):
    # H146-151: names only pypdf recovers survive via the union text layer.
    path = tmp_path / "catalogue.pdf"
    path.write_bytes(b"%PDF-fake")
    monkeypatch.setattr("pymupdf4llm.to_markdown", lambda p: "AirSense 11 overview")

    class _Page:
        def extract_text(self):
            return "Ultra Fine Filter FX2"

    monkeypatch.setattr("pypdf.PdfReader", lambda p: type("R", (), {"pages": [_Page()]})())
    doc = read_document(path, parser_union=True)
    assert "AirSense 11 overview" in doc.text
    assert "Ultra Fine Filter FX2" in doc.text
    # union off keeps only the primary parser
    off = read_document(path, parser_union=False)
    assert "Ultra Fine Filter FX2" not in off.text


def test_pdf_glyph_normalization_strips_trademark(tmp_path, monkeypatch):
    # H190: parser post-process cleans trademark glyphs so names are recoverable.
    path = tmp_path / "spec.pdf"
    path.write_bytes(b"%PDF-fake")
    monkeypatch.setattr("pymupdf4llm.to_markdown", lambda p: "SleepStyle™ Auto")
    doc = read_document(path, parser_union=False, glyph_normalization=True)
    assert doc.text == "SleepStyle Auto"


def test_document_id_deterministic(tmp_path):
    path = tmp_path / "same.md"
    path.write_text("x")
    assert read_document(path).id == read_document(path).id


def test_empty_document_text(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("")
    doc = read_document(path)
    assert doc.text == ""


def test_unknown_extension_raises(tmp_path):
    path = tmp_path / "data.xyz"
    path.write_text("?")
    with pytest.raises(UnsupportedFormatError) as exc:
        read_document(path)
    for ext in sorted(STRUCTURED_EXTENSIONS | UNSTRUCTURED_EXTENSIONS):
        assert ext in str(exc.value)


def test_read_structured_unknown_extension_raises(tmp_path):
    path = tmp_path / "data.xyz"
    path.write_text("?")
    with pytest.raises(UnsupportedFormatError):
        read_structured(path)


# -- structured -------------------------------------------------------------


def test_read_csv(tmp_path):
    path = tmp_path / "t.csv"
    path.write_text("name,price\nA,1.5\nB,2.0\n")
    records = read_structured(path)
    assert records == [{"name": "A", "price": 1.5}, {"name": "B", "price": 2.0}]


def test_read_tsv(tmp_path):
    path = tmp_path / "t.tsv"
    path.write_text("name\tprice\nA\t1.5\n")
    records = read_structured(path)
    assert records == [{"name": "A", "price": 1.5}]


def test_read_csv_nan_to_none(tmp_path):
    path = tmp_path / "t.csv"
    path.write_text("name,price\nA,\nB,2.0\n")
    records = read_structured(path)
    assert records[0]["price"] is None
    assert records[1]["price"] == 2.0


def test_read_json_list(tmp_path):
    path = tmp_path / "t.json"
    path.write_text('[{"a": 1}, {"a": 2}]')
    assert read_structured(path) == [{"a": 1}, {"a": 2}]


def test_read_json_single_object(tmp_path):
    path = tmp_path / "t.json"
    path.write_text('{"a": 1}')
    assert read_structured(path) == [{"a": 1}]


def test_read_jsonl(tmp_path):
    path = tmp_path / "t.jsonl"
    path.write_text('{"a": 1}\n{"a": 2}\n')
    assert read_structured(path) == [{"a": 1}, {"a": 2}]


def test_read_xlsx(tmp_path):
    path = tmp_path / "t.xlsx"
    pd.DataFrame({"name": ["A"], "qty": [3]}).to_excel(path, index=False)
    records = read_structured(path)
    assert records == [{"name": "A", "qty": 3}]


def test_read_parquet(tmp_path):
    path = tmp_path / "t.parquet"
    pd.DataFrame({"name": ["A", "B"], "qty": [1, 2]}).to_parquet(path)
    records = read_structured(path)
    assert records == [{"name": "A", "qty": 1}, {"name": "B", "qty": 2}]


# -- classification ---------------------------------------------------------


def test_is_structured_and_supported():
    assert is_structured(Path("x.csv"))
    assert not is_structured(Path("x.md"))
    assert is_supported(Path("x.pdf"))
    assert is_supported(Path("x.jsonl"))
    assert not is_supported(Path("x.xyz"))
    assert ".parquet" in STRUCTURED_EXTENSIONS
    assert ".md" in UNSTRUCTURED_EXTENSIONS


# -- expansion --------------------------------------------------------------


def test_iter_single_file(tmp_path):
    path = tmp_path / "one.md"
    path.write_text("x")
    assert iter_source_files(path) == [path]


def test_iter_directory_recursive_sorted(tmp_path):
    (tmp_path / "sub").mkdir()
    (tmp_path / "b.md").write_text("b")
    (tmp_path / "a.txt").write_text("a")
    (tmp_path / "sub" / "c.csv").write_text("x\n1\n")
    (tmp_path / "skip.xyz").write_text("?")
    files = iter_source_files(tmp_path)
    assert files == sorted(files)
    assert [f.name for f in files] == ["a.txt", "b.md", "c.csv"]


def test_iter_zip_expansion_sorted(tmp_path):
    zip_path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(zip_path, "w") as zf:
        zf.writestr("z.md", "z")
        zf.writestr("a.txt", "a")
        zf.writestr("nested/m.csv", "x\n1\n")
        zf.writestr("ignore.xyz", "?")
    workdir = tmp_path / "work"
    workdir.mkdir()
    files = iter_source_files(zip_path, workdir=workdir)
    assert files == sorted(files)
    assert [f.name for f in files] == ["a.txt", "m.csv", "z.md"]
    assert all(workdir in f.parents for f in files)
