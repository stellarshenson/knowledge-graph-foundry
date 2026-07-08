"""File readers for structured and unstructured sources.

Readers are pure: they parse files into Documents or records and raise on
failure - callers decide how to handle errors. Dispatch is by extension.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import tempfile
import zipfile

from loguru import logger

from knowledge_graph_foundry.models import Document, glyph_clean_text

STRUCTURED_EXTENSIONS = frozenset({".parquet", ".csv", ".tsv", ".xlsx", ".json", ".jsonl"})
UNSTRUCTURED_EXTENSIONS = frozenset({".pdf", ".docx", ".html", ".htm", ".md", ".txt"})


class UnsupportedFormatError(ValueError):
    """Raised when a file extension has no registered reader."""

    def __init__(self, path: Path) -> None:
        supported = ", ".join(sorted(STRUCTURED_EXTENSIONS | UNSTRUCTURED_EXTENSIONS))
        super().__init__(
            f"unsupported format '{path.suffix}' for {path.name}; supported: {supported}"
        )


def is_structured(path: Path) -> bool:
    """True when the extension maps to a structured (tabular/record) reader."""
    return path.suffix.lower() in STRUCTURED_EXTENSIONS


def is_supported(path: Path) -> bool:
    """True when any reader handles the extension."""
    ext = path.suffix.lower()
    return ext in STRUCTURED_EXTENSIONS or ext in UNSTRUCTURED_EXTENSIONS


def _document_id(path: Path) -> str:
    return "d_" + hashlib.sha1(path.name.encode()).hexdigest()[:16]


def read_document(
    path: Path, parser_union: bool = True, glyph_normalization: bool = True
) -> Document:
    """Read an unstructured file into a Document with markdown-ish text.

    R15-H146-151: `parser_union` adds pypdf's text layer to pymupdf4llm's for
    PDFs (union recovers 93.7% of names the project parser loses). R15-H190:
    `glyph_normalization` cleans trademark/unicode glyphs from the parsed text.
    """
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pymupdf4llm

        text = pymupdf4llm.to_markdown(str(path))
        if parser_union:
            text = _union_pypdf(path, text)
    elif ext == ".docx":
        text = _read_docx(path)
    elif ext in (".html", ".htm"):
        from bs4 import BeautifulSoup

        text = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser").get_text(
            separator="\n"
        )
    elif ext in (".md", ".txt"):
        text = path.read_text(encoding="utf-8")
    else:
        raise UnsupportedFormatError(path)

    if glyph_normalization:
        text = glyph_clean_text(text)

    logger.debug("read {} ({} chars)", path.name, len(text))
    return Document(
        id=_document_id(path),
        path=str(path),
        format=ext.lstrip("."),
        text=text,
        metadata={"file_name": path.name, "file_size": path.stat().st_size},
    )


def _union_pypdf(path: Path, base_text: str) -> str:
    """Append pypdf's extracted text layer to the primary parse (H146-151 union).

    pypdf recovers entity-name strings pymupdf4llm drops in table-heavy PDFs;
    concatenation is enough for name-presence recovery. A failing partner (e.g. a
    malformed PDF) is tolerated - the primary parse still stands."""
    from pypdf import PdfReader

    try:
        pages = PdfReader(str(path)).pages
        extra = "\n".join((page.extract_text() or "") for page in pages)
    except Exception as exc:  # noqa: BLE001 - union partner is best-effort
        logger.debug("pypdf union partner failed for {}: {}", path.name, exc)
        return base_text
    if not extra.strip():
        return base_text
    return base_text + "\n" + extra


def _read_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def read_structured(path: Path) -> list[dict]:
    """Read a structured file into a list of records with columns preserved."""
    import pandas as pd

    ext = path.suffix.lower()
    if ext == ".parquet":
        df = pd.read_parquet(path)
    elif ext == ".csv":
        df = pd.read_csv(path)
    elif ext == ".tsv":
        df = pd.read_csv(path, sep="\t")
    elif ext == ".xlsx":
        df = pd.read_excel(path, engine="openpyxl")
    elif ext == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, list) else [data]
    elif ext == ".jsonl":
        return [
            json.loads(line)
            for line in path.read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
    else:
        raise UnsupportedFormatError(path)

    df = df.astype(object).where(df.notna(), None)
    records = df.to_dict(orient="records")
    logger.debug("read {} ({} records)", path.name, len(records))
    return records


def iter_source_files(path: Path, workdir: Path | None = None) -> list[Path]:
    """Expand an input path to a deterministic (sorted) list of source files.

    A single file returns [file]; a directory returns its supported files
    recursively; a zip extracts to workdir (or a temp dir) first.
    """
    if path.is_dir():
        return sorted(p for p in path.rglob("*") if p.is_file() and is_supported(p))
    if path.suffix.lower() == ".zip":
        target = workdir or Path(tempfile.mkdtemp(prefix="kgf_ingest_"))
        with zipfile.ZipFile(path) as zf:
            zf.extractall(target)
        return sorted(p for p in target.rglob("*") if p.is_file() and is_supported(p))
    return [path]
